"""Local-first rolling Field Wise history and conservative material comparison.

Imported ticket snapshots are retained. Candidate links are never proof of billing.
"""
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date,datetime,timedelta

SOLD_STATUSES = {'SUBMITTED','APPROVED','INVOICED'}


def normalized(value):
    if isinstance(value,(int,float)) and float(value).is_integer(): value=int(value)
    return re.sub(r'\s+',' ',str(value or '').strip()).upper()


def fingerprint(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()).hexdigest()


def init_schema(con):
    con.executescript('''
      CREATE TABLE IF NOT EXISTS material_sales_imports (
        id INTEGER PRIMARY KEY AUTOINCREMENT, file_hash TEXT NOT NULL,
        source_file TEXT NOT NULL, window_start TEXT NOT NULL, window_end TEXT NOT NULL,
        imported_at TEXT NOT NULL, imported_by TEXT NOT NULL, row_count INTEGER NOT NULL,
        UNIQUE(file_hash,window_start,window_end)
      );
      CREATE TABLE IF NOT EXISTS material_sales_tickets (
        ticket_number TEXT PRIMARY KEY, ticket_date TEXT NOT NULL,
        content_hash TEXT NOT NULL, rows_json TEXT NOT NULL,
        active INTEGER NOT NULL, revision INTEGER NOT NULL, last_import_id INTEGER NOT NULL
      );
      CREATE TABLE IF NOT EXISTS material_sales_ticket_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT, ticket_number TEXT NOT NULL,
        revision INTEGER NOT NULL, import_id INTEGER NOT NULL,
        active INTEGER NOT NULL, content_hash TEXT NOT NULL, rows_json TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS material_sales_coverage (
        day TEXT PRIMARY KEY, import_id INTEGER NOT NULL
      );
      CREATE TABLE IF NOT EXISTS material_purchase_parse_cache (
        cache_key TEXT PRIMARY KEY, invoice_file TEXT NOT NULL, result_json TEXT NOT NULL
      );
      CREATE TABLE IF NOT EXISTS material_confirmed_matches (
        id INTEGER PRIMARY KEY AUTOINCREMENT, purchase_key TEXT NOT NULL,
        sales_line_id TEXT NOT NULL, bought_qty REAL NOT NULL,
        bought_units_per_sale_unit REAL NOT NULL, unit_note TEXT NOT NULL,
        confirmed_at TEXT NOT NULL, confirmed_by TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1
      );
    ''')


def import_sales_rows(con, raw_rows, window_end, source_file, file_hash, actor):
    """Replace tickets in the stated complete seven-day window, not all history.

    Rows for a ticket are a multiset: duplicate identical sales lines remain distinct.
    Ticket number is the source identifier; a changed job/customer is a revision.
    """
    end=date.fromisoformat(window_end); start=end-timedelta(days=6)
    grouped=defaultdict(list)
    for raw in raw_rows:
        row=dict(raw)
        day=date.fromisoformat(row['ticket_date'])
        if not start<=day<=end: raise ValueError('The export contains rows outside the selected seven-day window.')
        ticket=normalized(row.get('ticket_number'))
        if not ticket: raise ValueError('Every sales row must have a ticket number.')
        row['ticket_number']=ticket
        row['order_number']=normalized(row.get('order_number'))
        for key in ('quantity','rate','amount'):
            try: row[key]=float(row.get(key) or 0)
            except (ValueError,TypeError): raise ValueError(f'Invalid {key} on ticket {ticket}.')
            if not math.isfinite(row[key]): raise ValueError(f'Invalid {key} on ticket {ticket}.')
        grouped[ticket].append(row)
    for ticket,rows in grouped.items():
        if len({r['ticket_date'] for r in rows})!=1: raise ValueError(f'Ticket {ticket} has conflicting dates.')
    old_import=con.execute('SELECT id FROM material_sales_imports WHERE file_hash=? AND window_start=? AND window_end=?',(file_hash,start.isoformat(),end.isoformat())).fetchone()
    if old_import: return {'duplicate_file':True,'import_id':old_import['id'],'changed_tickets':0,'removed_tickets':0,'row_count':len(raw_rows)}
    latest=con.execute('SELECT MAX(window_end) FROM material_sales_imports').fetchone()[0]
    if latest and end.isoformat()<latest:
        raise ValueError('This export is older than the latest loaded window. Historical backfill must be reviewed separately so it cannot overwrite newer ticket corrections.')
    batch=con.execute('''INSERT INTO material_sales_imports (file_hash,source_file,window_start,window_end,imported_at,imported_by,row_count) VALUES (?,?,?,?,?,?,?)''',
        (file_hash,source_file,start.isoformat(),end.isoformat(),datetime.now().isoformat(timespec='seconds'),actor,len(raw_rows))).lastrowid
    changed=removed=0
    existing={r['ticket_number']:dict(r) for r in con.execute('SELECT * FROM material_sales_tickets')}
    def persist(ticket,rows,active,old):
        payload=json.dumps(sorted(rows,key=lambda r:json.dumps(r,sort_keys=True)),sort_keys=True)
        digest=fingerprint(json.loads(payload)); rev=(old['revision'] if old else 0)+1
        con.execute('''INSERT INTO material_sales_tickets VALUES (?,?,?,?,?,?,?) ON CONFLICT(ticket_number) DO UPDATE SET ticket_date=excluded.ticket_date,content_hash=excluded.content_hash,rows_json=excluded.rows_json,active=excluded.active,revision=excluded.revision,last_import_id=excluded.last_import_id''',
          (ticket,rows[0]['ticket_date'],digest,payload,active,rev,batch))
        con.execute('INSERT INTO material_sales_ticket_history (ticket_number,revision,import_id,active,content_hash,rows_json) VALUES (?,?,?,?,?,?)',(ticket,rev,batch,active,digest,payload))
    for ticket,rows in grouped.items():
        old=existing.get(ticket)
        canonical=sorted(rows,key=lambda r:json.dumps(r,sort_keys=True))
        if old and old['active'] and old['content_hash']==fingerprint(canonical): continue
        persist(ticket,rows,1,old); changed+=1
    for ticket,old in existing.items():
        if old['active'] and start.isoformat()<=old['ticket_date']<=end.isoformat() and ticket not in grouped:
            persist(ticket,json.loads(old['rows_json']),0,old);removed+=1
    for offset in range(7):
        con.execute('INSERT INTO material_sales_coverage VALUES (?,?) ON CONFLICT(day) DO UPDATE SET import_id=excluded.import_id',((start+timedelta(days=offset)).isoformat(),batch))
    return {'duplicate_file':False,'import_id':batch,'changed_tickets':changed,'removed_tickets':removed,'row_count':len(raw_rows)}


def read_fieldwise_export(path):
    import openpyxl
    workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        sheet=workbook['LineItems'] if 'LineItems' in workbook.sheetnames else workbook.worksheets[0]
        iterator=sheet.iter_rows(values_only=True)
        headers=[str(v or '').strip().lower() for v in next(iterator)]
        required=['ticket number','ticket date','order number','customer name','item','description','quantity','rate','sub total','status']
        missing=set(required)-set(headers)
        if missing: raise ValueError('Missing columns: '+', '.join(sorted(missing)))
        result=[]
        for values in iterator:
            if not any(values): continue
            data=dict(zip(headers,values)); d=data['ticket date']
            if isinstance(d,datetime): d=d.date()
            elif not isinstance(d,date):
                try: d=datetime.fromisoformat(str(d)).date()
                except ValueError: raise ValueError('Unrecognized ticket date: '+str(d))
            result.append({'ticket_number':data['ticket number'],'ticket_date':d.isoformat(),
              'order_number':data['order number'],'customer':str(data['customer name'] or ''),
              'item':str(data['item'] or ''),'description':str(data['description'] or ''),
              'quantity':data['quantity'],'rate':data['rate'],'amount':data['sub total'],
              'status':str(data['status'] or '')})
        return result
    finally: workbook.close()


def active_sales_rows(con):
    result=[]
    for ticket in con.execute('SELECT * FROM material_sales_tickets WHERE active=1 ORDER BY ticket_date,ticket_number'):
        counts=Counter()
        for row in json.loads(ticket['rows_json']):
            # Submitted -> Approved -> Invoiced does not change the underlying line.
            # A status change outside the allowed set still invalidates confirmation.
            digest=fingerprint({k:v for k,v in row.items() if k!='status'});counts[digest]+=1
            result.append({**row,'sales_line_id':ticket['ticket_number']+':'+digest[:20]+':'+str(counts[digest]),'ticket_revision':ticket['revision'],'counts_as_sold':normalized(row.get('status')) in SOLD_STATUSES})
    return result


def compare_candidates(invoices, documents, sales, catalog, coverage):
    """Review suggestions only. Never allocate one sales quantity to multiple buys."""
    by_code=defaultdict(list)
    catalog_labels={normalized(c['tablet_label']):c for c in catalog}
    for sale in sales:
        part=catalog_labels.get(normalized(sale['item']))
        if part: by_code[normalized(part['part_number'])].append(sale)
    rows=[];document_counts=Counter(i.get('sha256') for i in invoices)
    for invoice in invoices:
        doc=documents.get(invoice.get('sha256'),{})
        items=doc.get('items',[])
        if not items:
            rows.append({'invoice_id':invoice['id'],'po_number':invoice['po_number'],'invoice_file':invoice['invoice_file'],'job':invoice['job_number'],'vendor':invoice['vendor'],
              'invoice_date':doc.get('date'),'uploaded_at':invoice.get('uploaded_at'),'state':'Needs invoice parsing','description':doc.get('parser_error') or 'No reliable item lines extracted.','bought_qty':None,'candidates':[]})
            continue
        for item in items:
            code=normalized(item['item'])
            sku=code.replace(' ','')
            if normalized(doc.get('vendor'))=='DAKOTA SUPPLY GROUP':
                # DSG online invoices print their catalogue number explicitly as Pn:
                # The existing parser's item field can contain only the manufacturer.
                numbers=set(re.findall(r'\bPn:\s*(\d+)\b',item.get('description',''),flags=re.I))
                if len(numbers)==1: sku='DSG'+next(iter(numbers))
                elif sku.isdigit(): sku='DSG'+sku
            candidates=by_code.get(sku,[])
            same=[s for s in candidates if normalized(s['order_number'])==normalized(invoice['job_number']) and normalized(s['customer'])==normalized(invoice.get('customer_name'))]
            state='No exact item link in available sales'
            if code in ('SALES TAX','SHIPPING','FREIGHT','TARIFF','LOGISTIC SERVICE','SUPPLY CHAIN SURCHARGE'): state='Tax / freight / fee — separate review'
            elif float(item.get('amount') or 0)<0 or float(invoice.get('invoice_amount') or 0)<0: state='Credit / return — review'
            elif document_counts[invoice.get('sha256')]>1: state='Duplicate document attachment — review'
            elif doc.get('parse_status')!='Reconciled lines': state='Invoice lines do not reconcile — review'
            elif same and any(s['counts_as_sold'] and (float(s['quantity'])<0 or float(s['amount'])<0) for s in same): state='Ticket credit / return — review'
            elif same: state='Candidate sale — verify units and quantity' if any(s['counts_as_sold'] and float(s['quantity'])>0 for s in same) else 'Listed on ticket — status does not count as sold'
            elif candidates: state='Possible different job/customer — review'
            else:
                try: invoice_day=date.fromisoformat(doc.get('date') or '').isoformat()
                except ValueError: invoice_day=None
                if not invoice_day: state='Needs invoice date review'
                elif not coverage or invoice_day<min(coverage) or invoice_day>max(coverage): state='Outside accumulated sales coverage'
            rows.append({'purchase_key':fingerprint([invoice['id'],invoice.get('sha256'),item]),'invoice_id':invoice['id'],'po_number':invoice['po_number'],'invoice_file':invoice['invoice_file'],'job':invoice['job_number'],'customer':invoice.get('customer_name') or '', 'vendor':invoice['vendor'],'invoice_number':doc.get('invoice_number'),'invoice_date':doc.get('date'),
              'uploaded_at':invoice.get('uploaded_at'),'line_number':item['line_number'],'part_code':item['item'],'description':item['description'],'bought_qty':item['qty'],'unit':item.get('unit',''),'purchase_unit_cost':item['unit_cost'],'purchase_line_amount':item['amount'],
              'state':state,'candidates':same or candidates,'note':'Suggestions only. Sales quantities have not been allocated or confirmed. No billing conclusion from job number alone.'})
    return rows


def valid_matches(con, purchases, sales):
    buy_keys={p.get('purchase_key') for p in purchases if p.get('base_state',p.get('state'))=='Candidate sale — verify units and quantity'}
    sale_ids={s['sales_line_id'] for s in sales if s['counts_as_sold'] and float(s['quantity'])>0}
    current=[];stale=[]
    for row in con.execute('SELECT * FROM material_confirmed_matches WHERE active=1'):
        (current if row['purchase_key'] in buy_keys and row['sales_line_id'] in sale_ids else stale).append(dict(row))
    return current,stale


def in_forward_workflow(purchase, start_date):
    """Use purchase date for scope; new undated uploads need date review.

    Upload date admits an undated document for review, never substitutes for
    its purchase date or resets the grace period of a late uploaded invoice.
    """
    start=date.fromisoformat(start_date)
    try:return date.fromisoformat(purchase.get('invoice_date') or '')>=start
    except (ValueError,TypeError):
        try:return date.fromisoformat(str(purchase.get('uploaded_at') or '')[:10])>=start
        except ValueError:return False


def purchase_timing(purchase, coverage, today):
    """Age in calendar days through the latest fully loaded sales day.

    Review is not a finding of missed billing: quantities may be unconfirmed.
    Require every day since purchase, including the grace period, to be covered.
    """
    covered={date.fromisoformat(d) for d in coverage if d<=today.isoformat()}
    as_of=max(covered) if covered else None
    result={'timing_state':'Insufficient sales history','age_days':None,
            'timing_as_of':as_of.isoformat() if as_of else None}
    try: bought=date.fromisoformat(purchase.get('invoice_date') or '')
    except (ValueError,TypeError):
        result['timing_state']='Cannot determine — invoice date missing';return result
    if not as_of:return result
    age=(as_of-bought).days;result['age_days']=age
    qty=purchase.get('bought_qty')
    if qty is not None and float(qty)>0 and purchase.get('confirmed_qty',0)>=float(qty)-1e-8:
        result['timing_state']='Confirmed sold';return result
    if age<0:
        result['timing_state']='Purchase after loaded sales date';return result
    if any(bought+timedelta(days=n) not in covered for n in range(age+1)):
        return result
    if age<=3:
        result['timing_state']='Within 0–3 days — waiting for sale'
    elif qty is None or purchase.get('base_state',purchase.get('state')) in {
        'Needs invoice parsing','Invoice lines do not reconcile — review',
        'Duplicate document attachment — review','Credit / return — review',
        'Tax / freight / fee — separate review'}:
        result['timing_state']='Over 3 days — resolve source issue'
    else:
        result['timing_state']='Over 3 days — review unconfirmed quantity'
    return result


def confirm_match(con,purchases,sales,data,actor):
    buy=next((b for b in purchases if b.get('purchase_key')==data.get('purchase_key')),None)
    sale=next((s for s in sales if s['sales_line_id']==data.get('sales_line_id')),None)
    if not buy or not sale: raise ValueError('The source line changed or is no longer available. Refresh before confirming.')
    if buy['state']!='Candidate sale — verify units and quantity': raise ValueError('Resolve the purchase parsing, job or identity issue before confirming a sale.')
    if sale['sales_line_id'] not in {s['sales_line_id'] for s in buy['candidates']} or not sale['counts_as_sold']:
        raise ValueError('Choose a matching sale in Submitted, Approved, or Invoiced status.')
    try:
        qty=float(data['bought_qty']);factor=float(data['bought_units_per_sale_unit'])
        if not all(math.isfinite(n) and n>0 for n in (qty,factor)): raise ValueError()
    except (TypeError,KeyError,ValueError): raise ValueError('Enter a positive quantity and a verified unit conversion.')
    note=str(data.get('unit_note') or '').strip()
    if not note: raise ValueError('Describe the verified units, for example: 1 EA purchased = 1 EA on ticket.')
    matches,_=valid_matches(con,purchases,sales)
    buy_used=sum(m['bought_qty'] for m in matches if m['purchase_key']==buy['purchase_key'])
    sale_used=sum(m['bought_qty']/m['bought_units_per_sale_unit'] for m in matches if m['sales_line_id']==sale['sales_line_id'])
    if buy_used+qty>float(buy['bought_qty'])+1e-8: raise ValueError('This exceeds the remaining purchased quantity.')
    if sale_used+qty/factor>float(sale['quantity'])+1e-8: raise ValueError('That ticket quantity has already been used or is too small. It cannot cover two purchases twice.')
    return con.execute('''INSERT INTO material_confirmed_matches (purchase_key,sales_line_id,bought_qty,bought_units_per_sale_unit,unit_note,confirmed_at,confirmed_by) VALUES (?,?,?,?,?,?,?)''',
        (buy['purchase_key'],sale['sales_line_id'],qty,factor,note,datetime.now().isoformat(timespec='seconds'),actor)).lastrowid
